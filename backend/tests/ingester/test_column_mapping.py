
from agrosense.adapters.ingester.column_mapping import (
    MAPPING_VERSION,
    is_blank,
    map_columns,
    parse_alive,
)


def test_maps_real_annex_columns():
    cols = [
        "ID_MUEST", "ID Parcela", "LOCALIDAD", "Especie_M1", "Familia", "NombCom_M1",
        "Gremio ecológico de la especie",
        "Altura", "Coord_X", "Coord_Y",
        "Altura total (m)_M1", "Altura total (m)_M2", "Altura total (m)_M3",
        "Altura total (m)_M4",
        "Diámetro. Copa (m)_M1", "Diámetro. Copa (m)_M2", "Diámetro. Copa (m)_M3",
        "Diámetro. Copa (m)_M4",
        "DAP (CM)", "DAP (CM) M3", "DAP M4",
        "Sobrevivemcia M1", "Sobrevivencia M2", "Sobrevivencia M3", "Sobrevivencia M4",
        "Estado Fitosanitario_M1", "Estado Fitosanitario_M2",
        "Estado Fitosanitario_M3", "Estado Fitosanitario_M4",
        "ColumnaRara", "OtraDesconocida",
    ]
    mapping = map_columns(cols)
    assert mapping["ID_MUEST"] == "tree_id"
    assert mapping["LOCALIDAD"] == "locality"
    assert mapping["Altura total (m)_M3"] == "height_m"
    assert mapping["Diámetro. Copa (m)_M2"] == "crown_diameter_m"
    assert mapping["DAP (CM) M3"] == "dap_cm"
    assert mapping["Sobrevivemcia M1"] == "alive"
    assert mapping["DAP (CM)"] == "dap_cm"
    assert mapping["DAP M4"] == "dap_cm"
    assert mapping["Estado Fitosanitario_M2"] == "phytosanitary"
    assert mapping["Gremio ecológico de la especie"] == "guild"
    assert mapping["Altura"] == "elevation_m"
    assert mapping["Especie_M1"] == "species"
    assert "ColumnaRara" not in mapping
    assert "OtraDesconocida" not in mapping


def test_campaigns_exposed():
    """El ingester necesita saber a que campana pertenece cada columna M{k}."""
    from agrosense.adapters.ingester.column_mapping import campaign_columns

    cols = campaign_columns(["Altura total (m)_M2", "Sobrevivemcia M1",
                             "Estado Fitosanitario_M4", "DAP (CM) M3"])
    # canonical -> {campaign: nombre_real}
    assert cols["height_m"] == {2: "Altura total (m)_M2"}
    assert cols["alive"][1] == "Sobrevivemcia M1"
    assert cols["phytosanitary"][4] == "Estado Fitosanitario_M4"
    assert cols["dap_cm"][3] == "DAP (CM) M3"


def test_mapping_version_changed_with_e0():
    """Cambiar el mapeo cambia la provenance (ADR-004): nueva version."""
    assert MAPPING_VERSION == "2026-09-21-e0"


def test_parse_alive_semantics():
    assert parse_alive("Vivo") is True
    assert parse_alive("Muerto") is False
    assert parse_alive(" ") is None
    assert parse_alive("NaN") is None
    assert parse_alive(float("nan")) is None
    assert parse_alive(None) is None
    assert parse_alive("") is None


def test_is_blank_helper():
    from agrosense.adapters.ingester.column_mapping import is_blank

    assert is_blank(None) is True
    assert is_blank(float("nan")) is True
    assert is_blank("") is True
    assert is_blank("   ") is True
    assert is_blank("Vivo") is False
    assert is_blank(0.0) is False


def test_nan_is_not_blank_zero():
    """0.0 es un valor REAL (marcador bajo_umbral), no un blanco."""
    assert is_blank(0.0) is False


# ── E0: columnas que antes se descartaban ──────────────────────────────────

# Las 42 cabeceras reales de la hoja Monitoreo_4, tal cual: dobles espacios,
# saltos de linea y la cabecera truncada "Colonizació" incluidos. Fijadas aqui
# para que el test no dependa de que el archivo de datos este en la maquina.
REAL_HEADERS = [
    "Proyecto", "Cobertura  donde se establecio el material vegetal",
    "Cobertura vegetal asociada", "LOCALIDAD", "Diseño floristico",
    "Unidad de monitoreo", "Codigo de unidad muestreo", "ID Parcela", "ID Ind.",
    "ID_MUEST", "Altura", "Coord_X", "Coord_Y", "Familia", "Especie_M1",
    "NombCom_M1", "Gremio ecológico de la especie", "Evento",
    "Altura total (m)_M1", "Altura total (m)_M2", "Altura total (m)_M3",
    "Altura total (m)_M4", "Diámetro. Copa (m)_M1", "Diámetro. Copa (m)_M2",
    "Diámetro. Copa (m)_M3", "Diámetro. Copa (m)_M4", "DAP (CM)", "DAP (CM) M3",
    "DAP M4", "Sobrevivemcia M1", "Sobrevivencia M2", "Sobrevivencia M3",
    "Sobrevivencia M4", "Estado Fitosanitario_M1", "Estado Fitosanitario_M2",
    "Estado Fitosanitario_M3", "Estado Fitosanitario_M4",
    "Colonizació\nEpífitas vasculares", "Colonizació\nEpífitas no vasculares",
    "Observa", "Responsables", "Anotador",
]


def test_plot_attributes_are_mapped():
    from agrosense.adapters.ingester.column_mapping import fixed_columns

    fixed = fixed_columns(REAL_HEADERS)
    assert fixed["sampling_unit_code"] == "Codigo de unidad muestreo"
    assert fixed["monitoring_unit"] == "Unidad de monitoreo"
    assert fixed["floristic_design"] == "Diseño floristico"
    assert fixed["associated_cover"] == "Cobertura vegetal asociada"
    # la cabecera real tiene DOBLE espacio y "establecio" sin tilde
    assert fixed["establishment_cover"] == (
        "Cobertura  donde se establecio el material vegetal"
    )
    assert fixed["field_notes"] == "Observa"


def test_file_level_columns_are_mapped():
    from agrosense.adapters.ingester.column_mapping import file_columns

    meta = file_columns(REAL_HEADERS)
    assert meta == {
        "project_label": "Proyecto",
        "event": "Evento",
        "field_crew": "Responsables",
        "recorder": "Anotador",
    }


def test_every_real_column_is_classified():
    """Ninguna columna del formato de campo queda sin decidir.

    Es el test que habria detectado el hallazgo H2: la ingesta descartaba en
    silencio nueve columnas, entre ellas la unidad de muestreo por la que se
    agrupa la validacion cruzada. Ahora una columna nueva o bien se mapea, o
    bien se ignora EXPLICITAMENTE y con motivo.
    """
    from agrosense.adapters.ingester.column_mapping import unclassified_columns

    assert unclassified_columns(REAL_HEADERS) == []


def test_every_ignored_column_has_a_reason():
    from agrosense.adapters.ingester.column_mapping import IGNORED_COLUMNS

    assert IGNORED_COLUMNS, "la lista de ignoradas no puede estar vacia en E0"
    for header, reason in IGNORED_COLUMNS.items():
        assert reason and len(reason) > 15, f"{header!r} ignorada sin motivo"


def test_unknown_column_is_reported_as_unclassified():
    from agrosense.adapters.ingester.column_mapping import unclassified_columns

    assert unclassified_columns(REAL_HEADERS + ["Columna Nueva"]) == ["Columna Nueva"]


def test_matching_tolerates_accents_spaces_and_case():
    """Cada proyecto llega con su propio Excel (decision D6: los proyectos
    varian). Una tilde o un espacio de mas no pueden hacer perder una columna."""
    from agrosense.adapters.ingester.column_mapping import fixed_columns

    variantes = [
        "Diseño florístico",  # con tilde
        "DISEÑO FLORISTICO",  # mayusculas
        "  Diseño   floristico ",  # espacios
    ]
    for v in variantes:
        assert fixed_columns([v]).get("floristic_design") == v, v


def test_monitorings_have_no_upper_bound():
    """E0: M5, M6... se reconocen igual que M1-M4."""
    from agrosense.adapters.ingester.column_mapping import campaign_columns

    cols = campaign_columns([
        "Altura total (m)_M5", "Sobrevivencia M7", "DAP (CM) M6",
        "Estado Fitosanitario_M12", "Diámetro. Copa (m)_M5",
    ])
    assert cols["height_m"] == {5: "Altura total (m)_M5"}
    assert cols["alive"] == {7: "Sobrevivencia M7"}
    assert cols["dap_cm"] == {6: "DAP (CM) M6"}
    assert cols["phytosanitary"] == {12: "Estado Fitosanitario_M12"}
    assert cols["crown_diameter_m"] == {5: "Diámetro. Copa (m)_M5"}


def test_parse_event_number():
    from agrosense.adapters.ingester.column_mapping import parse_event_number

    assert parse_event_number("Cuarto monitoreo") == 4
    assert parse_event_number("Primer monitoreo") == 1
    assert parse_event_number("primero") == 1
    assert parse_event_number("Tercer Monitoreo") == 3
    assert parse_event_number("Décimo monitoreo") == 10
    assert parse_event_number("Monitoreo 5") == 5
    assert parse_event_number("M3") == 3
    assert parse_event_number("Linea base") is None
    assert parse_event_number(None) is None


def test_real_file_header_is_fully_classified():
    """La prueba fuerte: la cabecera del archivo REAL, no una copia a mano.

    La lista REAL_HEADERS de arriba se copio mal la primera vez ("Colonización"
    en vez de la forma truncada "Colonizació" que trae el archivo) y solo este
    test lo detecto.
    """
    from pathlib import Path

    import pytest
    from openpyxl import load_workbook

    from agrosense.adapters.ingester.column_mapping import unclassified_columns

    ref = Path(__file__).parents[2] / "data" / "raw" / "anexo1.xlsx"
    if not ref.exists():
        pytest.skip("dataset de referencia local ausente")
    wb = load_workbook(ref, read_only=True)
    try:
        header = next(wb["Monitoreo_4"].iter_rows(max_row=1, values_only=True))
    finally:
        wb.close()
    assert unclassified_columns(list(header)) == []
