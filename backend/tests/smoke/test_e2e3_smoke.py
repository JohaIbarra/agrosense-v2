"""Smoke de E2/E3 contra Supabase REAL (`pytest -m supabase`).

Verifica en Postgres el camino completo de la epica 1 con datos crudos:
cargar -> fechar el monitoreo -> calcular el analisis -> guardarlo como
snapshot JSONB -> leerlo -> construir el .xlsx. Lo que SQLite no garantiza
y aqui si se prueba:

  - el `payload` viaja a JSONB y vuelve con la misma forma (listas, nulos,
    decimales), no como texto;
  - `UNIQUE(monitoring_id)` hace que recalcular REEMPLACE el snapshot en vez
    de acumular filas;
  - la cascada real borra los snapshots al borrar el proyecto.

Campana pequena a proposito: reingerir los 856 arboles del Anexo tardaria
minutos contra el pooler (ver #G). Cada test limpia lo que crea.
"""
from __future__ import annotations

import os
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text

from agrosense.adapters.analysis.engine import ExploratoryAnalysisEngine
from agrosense.adapters.db.models import MonitoringAnalysisRow, Project
from agrosense.adapters.db.repository import (
    CampaignRepository,
    ProjectAnalysisRepository,
    ProjectRepository,
)
from agrosense.adapters.report.xlsx import build_report
from agrosense.application.dtos import CampaignData, FileMetadata
from agrosense.application.use_cases.monitoring_analysis import (
    get_monitoring_analysis,
    get_report_data,
    update_monitoring,
)
from agrosense.application.use_cases.upload_campaign import upload_campaign
from agrosense.domain.entities import Observation, StatusSemantic, Tree
from tests.auth.keys import ENGINEER_A as OWNER

pytestmark = [
    pytest.mark.supabase,
    pytest.mark.skipif(
        not os.environ.get("DATABASE_URL"), reason="DATABASE_URL ausente (backend/.env)"
    ),
]

SMOKE_PROJECT = "smoke-e2e3-agrosense"
HOY = date(2026, 9, 22)


@pytest.fixture()
def session():
    from agrosense.adapters.db.session import get_session_factory

    s = get_session_factory()()
    yield s
    s.rollback()
    s.close()


@pytest.fixture()
def project(session):
    def _drop():
        p = session.query(Project).filter_by(name=SMOKE_PROJECT).first()
        if p:
            # Borrado por SQL: la cascada la hace Postgres, no el ORM (evita el
            # N+1 que tumbaba la conexion del pooler).
            session.execute(text("DELETE FROM projects WHERE id = :i"), {"i": p.id})
            session.commit()

    _drop()
    p = ProjectRepository(session).create(owner_id=OWNER, name=SMOKE_PROJECT)
    yield p
    session.rollback()
    _drop()


# --- datos crudos de campo (lo unico que sube el ingeniero) ---------------

_ESPECIES = ("Cedrela montana", "Dodonaea viscosa", "Senna viarum")


def _tree(i: int) -> Tree:
    return Tree(
        tree_id=f"SMOKE-E3-{i}",
        species=_ESPECIES[i % 3],
        family="Meliaceae" if i % 3 == 0 else None,
        locality="Guayabal" if i < 6 else "Tres Jotas",
        plot_id="11" if i < 6 else "12",
        sampling_unit_code="SMOKE/FR/11" if i < 6 else "SMOKE/FR/12",
        floristic_design="Rehabilitación vegetal",
        associated_cover="Bosque de galería",
    )


def _obs(i: int, m: int) -> Observation:
    """M1 todos vivos; en M2 muere uno de cada cinco y el resto crece."""
    muerto = m == 2 and i % 5 == 0
    return Observation(
        tree_id=f"SMOKE-E3-{i}",
        campaign=m,
        height_m=None if muerto else round(0.4 + 0.1 * i + 0.3 * (m - 1), 2),
        crown_diameter_m=None if muerto else 0.2 + 0.05 * i,
        dap_cm=None,
        dap_status=StatusSemantic.BAJO_UMBRAL_DAP,
        phytosanitary=None if muerto else ("Bueno" if i % 2 else "Regular"),
        alive=not muerto,
        colonization=None,
        field_notes=None,
    )


class _FakeSource:
    """Puerto CampaignSource: el smoke prueba la DB, no el parser de Excel."""

    def __init__(self, campaigns=(1, 2)):
        self._campaigns = campaigns

    def read(self, content: bytes, filename: str) -> CampaignData:
        return CampaignData(
            trees=[_tree(i) for i in range(10)],
            observations=[_obs(i, m) for m in self._campaigns for i in range(10)],
            mapping_version="smoke-e2e3",
            file_metadata=FileMetadata(project_label="SMOKE", field_crew="Cuadrilla smoke"),
        )


def _subir(session, project, campaigns=(1, 2), fecha=date(2024, 11, 15), content=b"x"):
    return upload_campaign(
        project_id=project.id,
        filename="smoke.xlsx",
        content=content,
        project_repo=ProjectRepository(session),
        campaign_repo=CampaignRepository(session),
        source=_FakeSource(campaigns),
        owner_id=OWNER,
        monitoring_date=fecha,
        today=HOY,
        analysis_repo=ProjectAnalysisRepository(session),
        engine=ExploratoryAnalysisEngine(),
    )


# --- tests ---------------------------------------------------------------


def test_la_carga_deja_un_snapshot_por_monitoreo(session, project):
    """Subir datos crudos analiza TODOS los monitoreos del archivo."""
    result = _subir(session, project)
    assert result.valid and result.analyzed == [1, 2]

    filas = session.scalars(
        select(MonitoringAnalysisRow).where(MonitoringAnalysisRow.project_id == project.id)
    ).all()
    assert len(filas) == 2
    assert {f.analysis_version for f in filas} == {ExploratoryAnalysisEngine.version}
    assert len({f.input_hash for f in filas}) == 1, "mismo dataset, mismo hash"


def test_el_payload_vuelve_de_jsonb_con_su_forma(session, project):
    """JSONB no es texto: listas, nulos y decimales vuelven como objetos."""
    _subir(session, project)
    session.expunge_all()

    dto = get_monitoring_analysis(
        project.id, 2, OWNER,
        ProjectRepository(session), ProjectAnalysisRepository(session),
        ExploratoryAnalysisEngine(),
    )
    payload = dto.payload
    assert isinstance(payload["sections"], list) and payload["sections"]
    supervivencia = next(s for s in payload["sections"] if s["id"] == "supervivencia")
    tabla = supervivencia["tables"][0]
    assert isinstance(tabla["columns"], list)
    valores = [fila for fila in tabla["rows"]]
    assert valores, "la tabla de supervivencia no puede venir vacia"
    # 8 de 10 vivos en M2 (mueren i=0 y i=5)
    proyecto_fila = [f for f in valores if f.get("property") in (None, "Proyecto", "Guayabal")]
    assert proyecto_fila
    assert any(isinstance(v, float) for f in valores for v in f.values() if v is not None)


def test_recalcular_reemplaza_el_snapshot_no_lo_duplica(session, project):
    """UNIQUE(monitoring_id): un archivo corregido pisa el analisis anterior."""
    _subir(session, project)
    antes = session.scalar(
        text(
            "SELECT computed_at FROM monitoring_analyses ma "
            "JOIN monitorings m ON m.id = ma.monitoring_id "
            "WHERE m.project_id = :p AND m.number = 1"
        ).bindparams(p=project.id)
    )
    # Un archivo corregido (otro sha256; el mismo lo rechaza la ingesta)
    _subir(session, project, content=b"y")
    filas = session.execute(
        text("SELECT COUNT(*) FROM monitoring_analyses WHERE project_id = :p"),
        {"p": project.id},
    ).scalar()
    assert filas == 2, "un snapshot por monitoreo, no uno por carga"
    despues = session.scalar(
        text(
            "SELECT computed_at FROM monitoring_analyses ma "
            "JOIN monitorings m ON m.id = ma.monitoring_id "
            "WHERE m.project_id = :p AND m.number = 1"
        ).bindparams(p=project.id)
    )
    assert despues >= antes


def test_la_fecha_del_monitoreo_se_guarda_como_date(session, project):
    """La fecha que escribe el ingeniero llega a la columna DATE de Postgres."""
    _subir(session, project, fecha=date(2024, 11, 15))
    valor = session.execute(
        text("SELECT monitoring_date FROM monitorings WHERE project_id = :p AND number = 2"),
        {"p": project.id},
    ).scalar()
    assert valor == date(2024, 11, 15)

    dto = update_monitoring(
        project.id, 1, OWNER, ProjectRepository(session), HOY,
        monitoring_date=date(2024, 5, 3), notes="Corrección de campo",
    )
    assert dto.monitoring_date == date(2024, 5, 3) and dto.notes == "Corrección de campo"


def test_el_reporte_sale_del_mismo_snapshot_que_la_pagina(session, project):
    """Una sola cifra de verdad: el .xlsx no recalcula nada."""
    _subir(session, project)
    data = get_report_data(
        project.id, 2, OWNER,
        ProjectRepository(session), ProjectAnalysisRepository(session),
        ExploratoryAnalysisEngine(),
    )
    blob = build_report(data)
    wb = load_workbook(BytesIO(blob))
    assert "Resumen" in wb.sheetnames and "Datos crudos" in wb.sheetnames
    # los datos crudos del reporte son los 10 arboles que subimos
    crudos = wb["Datos crudos"]
    assert crudos.max_row == 11, "cabecera + 10 arboles"


def test_borrar_el_proyecto_se_lleva_sus_snapshots(session, project):
    """La cascada es de la DB, no del ORM."""
    _subir(session, project)
    pid = project.id
    session.execute(text("DELETE FROM projects WHERE id = :i"), {"i": pid})
    session.commit()
    n = session.execute(
        text("SELECT COUNT(*) FROM monitoring_analyses WHERE project_id = :p"), {"p": pid}
    ).scalar()
    assert n == 0
