"""Reporte .xlsx (E3): una hoja por analisis + resumen + datos crudos, con las
MISMAS cifras del payload (no recalcula) y formato limpio."""

import io
from dataclasses import replace
from datetime import UTC, date, datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from agrosense.adapters.analysis.engine import ANALYSIS_VERSION, analyze_monitoring
from agrosense.adapters.ingester.ingest import ingest_wide
from agrosense.adapters.report.xlsx import build_report, report_filename
from agrosense.application.dtos import AnalysisDTO, ReportData
from tests.analysis.test_engine import OBS, TREES


def _data(number=2):
    payload = analyze_monitoring(TREES, OBS, number)
    return ReportData(
        project_name="Restauración Guayabal",
        project_code="AGS-2026-0001",
        analysis=AnalysisDTO(
            project_id=1,
            monitoring=number,
            monitoring_date=date(2025, 9, 1),
            analysis_version=ANALYSIS_VERSION,
            input_hash="a" * 64,
            computed_at=datetime(2026, 9, 22, 10, 0, tzinfo=UTC),
            payload=payload,
        ),
        trees=TREES,
        observations=[o for o in OBS if o.campaign <= number],
        monitoring_dates={1: date(2025, 3, 1), 2: date(2025, 9, 1)},
    )


@pytest.fixture(scope="module")
def workbook():
    data = _data()
    return data, load_workbook(io.BytesIO(build_report(data)))


def test_one_sheet_per_analysis_plus_summary_and_raw(workbook):
    _, wb = workbook
    assert wb.sheetnames == [
        "Resumen",
        "Composición",
        "Alturas",
        "Diámetro de copa",
        "Supervivencia",
        "Estado fitosanitario",
        "Edades",
        "DAP",
        "Datos crudos",
    ]


def _find(ws, text):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == text:
                return cell
    raise AssertionError(f"{text!r} no esta en {ws.title}")


def test_numbers_are_the_payload_numbers_with_their_format(workbook):
    data, wb = workbook
    ws = wb["Alturas"]
    table = next(
        t
        for s in data.analysis.payload["sections"]
        if s["id"] == "alturas"
        for t in s["tables"]
        if t["id"] == "alturas-predio-a-especie"
    )
    title = _find(ws, table["title"])
    header_row = title.row + 1
    headers = [c.value for c in ws[header_row]]
    assert headers[: len(table["columns"])] == [c["label"] for c in table["columns"]]
    first = ws[header_row + 1]
    cur = [c["key"] for c in table["columns"]].index("cur")
    assert first[cur].value == table["rows"][0]["cur"]
    assert first[cur].number_format == "0.00"


def test_percent_columns_keep_the_page_scale(workbook):
    data, wb = workbook
    ws = wb["Supervivencia"]
    cell = _find(ws, "% M2")
    # la primera fila de datos bajo el encabezado «% M2»
    value = ws.cell(row=cell.row + 1, column=cell.column).value
    assert 0 <= value <= 100
    assert ws.cell(row=cell.row + 1, column=cell.column).number_format == "0.0"


def test_charts_are_native_excel_charts(workbook):
    _, wb = workbook
    assert len(wb["Composición"]._charts) >= 1
    assert len(wb["Supervivencia"]._charts) >= 1


def test_summary_carries_provenance(workbook):
    _, wb = workbook
    ws = wb["Resumen"]
    assert _find(ws, "Versión del cálculo").offset(column=1).value == ANALYSIS_VERSION
    assert _find(ws, "Proyecto").offset(column=1).value == "Restauración Guayabal"


def test_raw_sheet_can_be_uploaded_again(workbook):
    data, wb = workbook
    buf = io.BytesIO()
    wb.save(buf)
    df = pd.read_excel(io.BytesIO(buf.getvalue()), sheet_name="Datos crudos")
    again = ingest_wide(df)
    assert sorted(t.tree_id for t in again.trees) == sorted(t.tree_id for t in data.trees)
    original = {(o.tree_id, o.campaign): o for o in data.observations}
    for o in again.observations:
        before = original[(o.tree_id, o.campaign)]
        assert (o.height_m, o.alive, o.dap_cm) == (before.height_m, before.alive, before.dap_cm)
    assert len(again.observations) == len(data.observations)


def test_column_widths_are_set(workbook):
    _, wb = workbook
    ws = wb["Datos crudos"]
    assert ws.column_dimensions["A"].width >= 8
    assert ws.freeze_panes == "B2"


def test_filename_is_safe():
    assert report_filename(_data()) == "AgroSense_AGS-2026-0001_M2.xlsx"


def test_field_text_never_becomes_an_excel_formula():
    """Un dato de campo que empieza por `=` es TEXTO, no una formula.

    openpyxl decide el tipo de celda mirando el valor: cualquier cadena que
    empiece por `=` queda marcada como formula y Excel la EJECUTA al abrir el
    reporte. Como la especie, el identificador y las notas salen del archivo
    que sube el usuario, eso convertiria el .xlsx en un vector de inyeccion
    de formulas. Regresion de la revision de seguridad de E2/E3.
    """
    base = _data()
    hostil = "=1+1"
    base.analysis.payload["sections"][0]["tables"][0]["rows"][0]["species"] = hostil
    data = replace(
        base,
        trees=[t.model_copy(update={"species": hostil}) for t in base.trees],
        observations=[
            o.model_copy(update={"field_notes": '=HYPERLINK("http://x")'})
            for o in base.observations
        ],
    )

    wb = load_workbook(io.BytesIO(build_report(data)))
    formulas = [
        (ws.title, cell.coordinate, cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert formulas == [], f"celdas ejecutables en el reporte: {formulas[:3]}"
    crudos = wb["Datos crudos"]
    assert _find(crudos, hostil).data_type == "s"
